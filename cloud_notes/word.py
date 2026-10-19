"""Editable Word notes with native OMML equations and local figures."""
from pathlib import Path
import io
import json
import os
import re
import tempfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
MATH_NS='http://schemas.openxmlformats.org/officeDocument/2006/math'
FENCE=chr(96)*3
def safe_name(value):
    text=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',str(value)).strip(' .')
    if not text or text.upper() in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}: text='课程_'+text
    return text[:65]
def math_run(paragraph,latex):
    from cloud_notes.equations import convert_equation
    paragraph._p.append(convert_equation(latex))
def inline(paragraph,text):
    text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text)
    text=text.replace(r'\(', '$').replace(r'\)', '$')
    for part in re.split(r'(\$[^$\n]+\$|\*\*[^*]+\*\*)',text):
        if part.startswith('$') and part.endswith('$') and len(part)>2: math_run(paragraph,part[1:-1])
        elif part.startswith('**') and part.endswith('**'): paragraph.add_run(part[2:-2]).bold=True
        else: paragraph.add_run(part.replace(chr(96),''))
def font(size=24):
    for name in ('C:/Windows/Fonts/msyh.ttc','/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'):
        if Path(name).exists(): return ImageFont.truetype(name,size)
    return ImageFont.load_default()
def diagram_image(spec):
    labels=spec.get('nodes',[])
    if not isinstance(labels,list) or not 2<=len(labels)<=6: raise ValueError('Invalid diagram')
    if any(not isinstance(s,str) or len(s)>80 for s in labels): raise ValueError('Invalid label')
    box_height=40+31*max((len(label)+29)//30 for label in labels)
    step=box_height+37
    image=Image.new('RGB',(1200,60+len(labels)*step),'white')
    draw=ImageDraw.Draw(image); f=font(28)
    for i,label in enumerate(labels):
        y=30+i*step
        draw.rounded_rectangle((110,y,1090,y+box_height),radius=16,fill='#edf4ff',outline='#315a92',width=2)
        for j,line in enumerate([label[k:k+30] for k in range(0,len(label),30)]):
            draw.text((135,y+16+j*31),line,font=f,fill='#152f50')
        if i<len(labels)-1:
            draw.line((600,y+box_height+1,600,y+box_height+34),fill='#315a92',width=3)
            draw.polygon([(593,y+box_height+24),(607,y+box_height+24),(600,y+box_height+34)],fill='#315a92')
    output=io.BytesIO(); image.save(output,format='PNG'); output.seek(0); return output
def create_word(row, root=ROOT, pictures=()):
    root=Path(root); folder=root/'documents'/safe_name(row['course_title'])
    folder.mkdir(parents=True,exist_ok=True)
    sid=safe_name(row['sub_id'])
    if len(sid)>32:
        import hashlib
        sid=sid[:24]+hashlib.sha256(str(row['sub_id']).encode()).hexdigest()[:8]
    path=folder/(safe_name(row.get('sub_title') or row.get('date') or sid)+'_'+sid+'.docx')
    if path.exists(): return path
    document=Document(); section=document.sections[0]
    section.top_margin=section.bottom_margin=Cm(2)
    section.left_margin=section.right_margin=Cm(2.2)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
        style=document.styles[name]; style.font.name='Microsoft YaHei'
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    document.styles['Normal'].font.size=Pt(11)
    document.styles['Normal'].paragraph_format.space_after=Pt(6)
    document.styles['Normal'].paragraph_format.line_spacing=1.3
    document.styles['Title'].font.size=Pt(24)
    document.add_paragraph(row['course_title'],style='Title')
    document.add_paragraph(row.get('sub_title') or row.get('date') or '课程笔记')
    document.add_paragraph('课程知识笔记与重点事项',style='Subtitle')
    text=row['summary'].replace(r'\[','$$').replace(r'\]','$$')
    pattern='('+re.escape(FENCE)+r'diagram\s*\n[\s\S]*?'+re.escape(FENCE)+r'|\$\$[\s\S]*?\$\$)'
    for block in re.split(pattern,text):
        if block.startswith(FENCE+'diagram'):
            try: spec=json.loads(block[len(FENCE+'diagram'): -len(FENCE)].strip())
            except (ValueError,TypeError): continue
            if not isinstance(spec,dict): continue
            document.add_paragraph(spec.get('title','知识关系示意图'),style='Heading 2')
            try: document.add_picture(diagram_image(spec),width=Cm(15.5))
            except (ValueError,TypeError):
                for label in spec.get('nodes',[])[:6]: document.add_paragraph(str(label),style='List Bullet')
            document.add_paragraph('根据本课笔记生成的解释性示意图，请结合正文理解。',style='Caption')
            continue
        if block.startswith('$$'):
            math_run(document.add_paragraph(),block[2:-2].strip()); continue
        lines=block.splitlines(); i=0
        while i<len(lines):
            line=lines[i].strip(); i+=1
            if not line or line.startswith(FENCE): continue
            if line.startswith('|') and i<len(lines) and re.match(r'^\s*\|?[\s:|-]+\|?\s*$',lines[i]):
                headers=[c.strip() for c in line.strip('|').split('|')]; i+=1
                table=document.add_table(rows=1,cols=len(headers)); table.style='Table Grid'
                for cell,title in zip(table.rows[0].cells,headers):
                    inline(cell.paragraphs[0],title)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F7');cell._tc.get_or_add_tcPr().append(shade)
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                while i<len(lines) and lines[i].strip().startswith('|'):
                    values=[c.strip() for c in lines[i].strip().strip('|').split('|')]; i+=1
                    for cell,value in zip(table.add_row().cells,values): inline(cell.paragraphs[0],value)
                continue
            heading=re.match(r'^(#{1,6})\s+(.*)',line)
            if heading:
                p=document.add_paragraph(style=f'Heading {min(3,max(1,len(heading.group(1))-2))}'); inline(p,heading.group(2))
            elif re.match(r'^[-*]\s+',line): inline(document.add_paragraph(style='List Bullet'),line[2:])
            elif re.match(r'^\d+\.\s+',line): inline(document.add_paragraph(style='List Number'),re.sub(r'^\d+\.\s+','',line))
            elif not re.match(r'^[-*_]{3,}$',line): inline(document.add_paragraph(),line)
    if pictures:
        document.add_paragraph('课堂图示参考',style='Heading 1')
        for label,picture in list(pictures)[:6]:
            try:
                document.add_picture(str(picture),width=Cm(15.5))
                document.add_paragraph(label,style='Caption')
            except (ValueError,OSError): continue
    fd,tmp=tempfile.mkstemp(suffix='.docx',dir=folder); os.close(fd)
    try: document.save(tmp); os.replace(tmp,path)
    finally:
        if os.path.exists(tmp): os.unlink(tmp)
    return path
def send_attachment(path,row,settings):
    import smtplib
    import ssl
    from email.message import EmailMessage
    message=EmailMessage()
    message['From']=settings['SMTP_EMAIL']; message['To']=settings['RECEIVER_EMAIL']
    message['Subject']=f"[iCourse] {row['course_title']} {row.get('sub_title','')} · Word 笔记"
    message.set_content('本次课程笔记已整理为 Word 文档，请查看附件。')
    message.add_attachment(Path(path).read_bytes(),maintype='application',
        subtype='vnd.openxmlformats-officedocument.wordprocessingml.document',filename=Path(path).name)
    with smtplib.SMTP_SSL('smtp.qq.com',465,context=ssl.create_default_context(),timeout=45) as smtp:
        smtp.login(settings['SMTP_EMAIL'],settings['SMTP_PASSWORD'])
        return not smtp.send_message(message)
